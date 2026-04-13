"""
Knowledge ingestion — read files, chunk them, store in ChromaDB.

Supported formats:
  .txt .md .py .js .ts .json .yaml .yml .csv .rst  (plain text)
  .pdf   requires: pip install pypdf
  .docx  requires: pip install python-docx
  .xlsx  requires: pip install openpyxl
"""
import os
from pathlib import Path
from typing import Optional

import chromadb
from chromadb.config import Settings

from .embeddings import OllamaEmbeddingFunction

_ollama_ef = OllamaEmbeddingFunction()

SUPPORTED_EXTENSIONS = {
    ".txt", ".md", ".py", ".js", ".ts", ".json",
    ".yaml", ".yml", ".csv", ".rst", ".html", ".xml",
    ".pdf", ".docx", ".xlsx",
}


# -----------------------------------------------------------------------
# ChromaDB client factory
# -----------------------------------------------------------------------
def get_chroma_client(persist_directory: str = "./data/chroma") -> chromadb.ClientAPI:
    os.makedirs(persist_directory, exist_ok=True)
    return chromadb.PersistentClient(
        path=persist_directory,
        settings=Settings(anonymized_telemetry=False),
    )


# -----------------------------------------------------------------------
# Chunking
# -----------------------------------------------------------------------
def chunk_text(text: str, chunk_size: int = 1000, overlap: int = 200) -> list[str]:
    """Split *text* into overlapping fixed-size chunks."""
    if len(text) <= chunk_size:
        return [text.strip()] if text.strip() else []
    chunks: list[str] = []
    start = 0
    while start < len(text):
        chunk = text[start : start + chunk_size].strip()
        if chunk:
            chunks.append(chunk)
        start += chunk_size - overlap
    return chunks


# -----------------------------------------------------------------------
# File readers
# -----------------------------------------------------------------------
def _read_pdf(path: str) -> str:
    from pypdf import PdfReader  # lazy import
    reader = PdfReader(path)
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def _read_docx(path: str) -> str:
    from docx import Document  # python-docx
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


def _read_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = []
    for sheet in wb.worksheets:
        lines.append(f"=== Sheet: {sheet.title} ===")
        for row in sheet.iter_rows(values_only=True):
            lines.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(lines)


def _read_file(path: str) -> str:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return _read_pdf(path)
    if ext == ".docx":
        return _read_docx(path)
    if ext == ".xlsx":
        return _read_xlsx(path)
    return Path(path).read_text(encoding="utf-8", errors="replace")


# -----------------------------------------------------------------------
# Public API
# -----------------------------------------------------------------------
async def ingest_file(
    file_path: str,
    collection_name: str = "default",
    persist_directory: str = "./data/chroma",
    chunk_size: int = 1000,
    overlap: int = 200,
) -> dict:
    """Ingest a single file into the named ChromaDB collection."""
    path = Path(file_path)
    if not path.exists():
        return {"success": False, "error": f"File not found: {file_path}"}
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        return {"success": False, "error": f"Unsupported extension: {path.suffix}"}

    try:
        content = _read_file(file_path)
    except ImportError as exc:
        return {"success": False, "error": f"Missing dependency: {exc}"}
    except Exception as exc:
        return {"success": False, "error": f"Read error: {exc}"}

    if not content.strip():
        return {"success": False, "error": "File is empty after parsing"}

    chunks = chunk_text(content, chunk_size=chunk_size, overlap=overlap)
    if not chunks:
        return {"success": False, "error": "No text chunks produced"}

    client = get_chroma_client(persist_directory)
    collection = client.get_or_create_collection(
        name=collection_name,
        metadata={"hnsw:space": "cosine"},
        embedding_function=_ollama_ef,
    )

    doc_ids = [f"{path.stem}_c{i}" for i in range(len(chunks))]
    metadatas = [
        {"source": str(path.resolve()), "filename": path.name, "chunk_index": i}
        for i in range(len(chunks))
    ]
    collection.upsert(ids=doc_ids, documents=chunks, metadatas=metadatas)

    return {
        "success": True,
        "file": str(path),
        "chunks": len(chunks),
        "collection": collection_name,
    }


async def ingest_directory(
    directory: str,
    collection_name: str = "default",
    persist_directory: str = "./data/chroma",
    recursive: bool = True,
    chunk_size: int = 1000,
    overlap: int = 200,
) -> dict:
    """Ingest all supported files in *directory*."""
    p = Path(directory)
    if not p.exists():
        return {"success": False, "error": f"Directory not found: {directory}"}

    pattern = "**/*" if recursive else "*"
    files = [
        f for f in p.glob(pattern)
        if f.is_file() and f.suffix.lower() in SUPPORTED_EXTENSIONS
    ]

    results = []
    for file_path in files:
        result = await ingest_file(
            str(file_path), collection_name, persist_directory, chunk_size, overlap
        )
        results.append(result)

    successful = sum(1 for r in results if r.get("success"))
    return {
        "success": True,
        "total": len(files),
        "successful": successful,
        "failed": len(files) - successful,
        "results": results,
    }
